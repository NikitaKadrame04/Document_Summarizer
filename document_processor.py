"""In-memory extraction and lightweight cleaning for PDF, DOCX, and TXT."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath
from typing import Any

from config import MAX_UPLOAD_BYTES

MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class DocumentProcessingError(ValueError):
    """An upload could not be safely or usefully converted to text."""


@dataclass(frozen=True)
class DocumentMetadata:
    filename: str
    file_type: str
    character_count: int
    word_count: int
    page_count: int | None = None
    low_text_pages: tuple[int, ...] = ()
    encoding: str | None = None


@dataclass(frozen=True)
class DocumentExtractionResult:
    text: str
    original_text: str
    metadata: DocumentMetadata


def clean_text(text: str) -> str:
    """Normalize line endings and controls while retaining readable layout."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "".join(
        char for char in text
        if char in "\n\t" or ord(char) >= 32 and ord(char) != 127
    )
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_uploaded_bytes(uploaded_file: Any) -> bytes:
    if isinstance(uploaded_file, (bytes, bytearray)):
        return bytes(uploaded_file)
    if hasattr(uploaded_file, "getvalue"):
        return uploaded_file.getvalue()
    if hasattr(uploaded_file, "read"):
        if hasattr(uploaded_file, "seek"):
            uploaded_file.seek(0)
        data = uploaded_file.read()
        if hasattr(uploaded_file, "seek"):
            uploaded_file.seek(0)
        return data
    raise DocumentProcessingError("The uploaded file could not be read.")


def _paragraph_and_table_text(doc: Any) -> str:
    """Read DOCX paragraphs and tables in document order when supported."""
    parts: list[str] = []
    if hasattr(doc, "iter_inner_content"):
        for item in doc.iter_inner_content():
            if hasattr(item, "rows"):
                for row in item.rows:
                    cells = [clean_text(cell.text) for cell in row.cells]
                    row_text = " | ".join(cell for cell in cells if cell)
                    if row_text:
                        parts.append(row_text)
            else:
                paragraph = clean_text(getattr(item, "text", ""))
                if paragraph:
                    parts.append(paragraph)
    else:  # Compatibility with older python-docx releases.
        parts.extend(clean_text(p.text) for p in doc.paragraphs if clean_text(p.text))
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(clean_text(cell.text) for cell in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)
    return "\n\n".join(parts)


def extract_document_text(uploaded_file: Any, file_type: str | None = None) -> DocumentExtractionResult:
    """Validate an upload and return normalized text plus file metadata.

    The upload is processed in memory; this function never writes it to disk.
    """
    filename = str(getattr(uploaded_file, "name", "uploaded_document"))
    extension = (file_type or PurePath(filename).suffix).lower().lstrip(".")
    if extension.startswith("."):
        extension = extension[1:]
    if extension not in {"pdf", "docx", "txt"}:
        raise DocumentProcessingError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")

    data = _read_uploaded_bytes(uploaded_file)
    if not data:
        raise DocumentProcessingError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        upload_limit_mb = MAX_UPLOAD_BYTES / (1024 * 1024)
        raise DocumentProcessingError(f"The file exceeds the configured {upload_limit_mb:g} MB upload limit.")

    page_count: int | None = None
    low_text_pages: tuple[int, ...] = ()
    encoding: str | None = None

    try:
        if extension == "txt":
            try:
                original = data.decode("utf-8-sig")
                encoding = "UTF-8"
            except UnicodeDecodeError:
                try:
                    original = data.decode("cp1252")
                    encoding = "Windows-1252"
                except UnicodeDecodeError as exc:
                    raise DocumentProcessingError("The text file could not be decoded as UTF-8 or Windows-1252.") from exc
        elif extension == "docx":
            if not zipfile.is_zipfile(io.BytesIO(data)):
                raise DocumentProcessingError("This DOCX file is malformed or is not a valid Word document.")
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise DocumentProcessingError("This file is not a valid DOCX document.")
                expanded_size = sum(member.file_size for member in archive.infolist())
                if expanded_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DocumentProcessingError("This DOCX expands beyond the safe processing limit.")
                bad_member = archive.testzip()
                if bad_member:
                    raise DocumentProcessingError("This DOCX file is damaged and cannot be read.")
            from docx import Document

            doc = Document(io.BytesIO(data))
            original = _paragraph_and_table_text(doc)
        else:
            if not data.startswith(b"%PDF-"):
                raise DocumentProcessingError("This file does not have a valid PDF header.")
            import fitz

            with fitz.open(stream=data, filetype="pdf") as pdf:
                if pdf.is_encrypted and pdf.authenticate("") <= 0:
                    raise DocumentProcessingError(
                        "This PDF is password-protected or encrypted. Remove its protection before uploading."
                    )
                page_count = pdf.page_count
                pages: list[str] = []
                low_pages: list[int] = []
                for index, page in enumerate(pdf, start=1):
                    page_text = clean_text(page.get_text("text"))
                    if len(page_text) < 20:
                        low_pages.append(index)
                    pages.append(page_text)
                original = "\n\n".join(pages)
                low_text_pages = tuple(low_pages)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        if extension == "pdf":
            raise DocumentProcessingError("The PDF is corrupted or could not be opened.") from exc
        if extension == "docx":
            raise DocumentProcessingError("The DOCX file is corrupted or could not be opened.") from exc
        raise DocumentProcessingError("The uploaded document could not be processed.") from exc

    text = clean_text(original)
    if not text:
        if extension == "pdf":
            raise DocumentProcessingError(
                "No selectable text was found in this PDF. It may be a scanned image; OCR is not included in this version."
            )
        raise DocumentProcessingError("No extractable text was found in this document.")

    metadata = DocumentMetadata(
        filename=filename,
        file_type=extension.upper(),
        character_count=len(text),
        word_count=len(text.split()),
        page_count=page_count,
        low_text_pages=low_text_pages,
        encoding=encoding,
    )
    return DocumentExtractionResult(text=text, original_text=original, metadata=metadata)
