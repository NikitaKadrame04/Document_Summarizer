from __future__ import annotations

import io

import pytest

import document_processor
from document_processor import DocumentProcessingError, clean_text, extract_document_text


class Upload:
    def __init__(self, name: str, content: bytes):
        self.name = name
        self._content = content

    def getvalue(self) -> bytes:
        return self._content


def test_txt_utf8_and_bom_metadata():
    result = extract_document_text(Upload("notes.txt", "\ufeffCafé\nSmart documents".encode("utf-8")))
    assert result.text == "Café\nSmart documents"
    assert result.metadata.encoding == "UTF-8"
    assert result.metadata.file_type == "TXT"
    assert result.metadata.word_count == 3


def test_txt_windows_1252_fallback():
    result = extract_document_text(Upload("legacy.txt", "Résumé".encode("cp1252")))
    assert result.text == "Résumé"
    assert result.metadata.encoding == "Windows-1252"


def test_empty_txt_rejected():
    with pytest.raises(DocumentProcessingError, match="empty"):
        extract_document_text(Upload("empty.txt", b""))
    with pytest.raises(DocumentProcessingError, match="No extractable text"):
        extract_document_text(Upload("blank.txt", b" \n\t"))


def test_docx_paragraphs_and_table():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Project overview")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "Evaluation"
    stream = io.BytesIO()
    doc.save(stream)

    result = extract_document_text(Upload("report.docx", stream.getvalue()))
    assert "Project overview" in result.text
    assert "Method | Evaluation" in result.text


def test_malformed_docx_rejected():
    with pytest.raises(DocumentProcessingError, match="malformed"):
        extract_document_text(Upload("bad.docx", b"not a zip file"))


def test_pdf_page_text_and_page_count():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A synthetic PDF page with extractable text.")
    data = pdf.tobytes()
    pdf.close()
    result = extract_document_text(Upload("sample.pdf", data))
    assert "synthetic PDF" in result.text
    assert result.metadata.page_count == 1


def test_empty_and_invalid_pdf_errors():
    with pytest.raises(DocumentProcessingError, match="empty"):
        extract_document_text(Upload("empty.pdf", b""))
    with pytest.raises(DocumentProcessingError, match="header"):
        extract_document_text(Upload("fake.pdf", b"not a PDF"))


def test_scanned_pdf_explains_ocr_limit():
    import fitz

    pdf = fitz.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    with pytest.raises(DocumentProcessingError, match="OCR is not included"):
        extract_document_text(Upload("scan.pdf", data))


def test_unsupported_extension_rejected():
    with pytest.raises(DocumentProcessingError, match="Unsupported"):
        extract_document_text(Upload("program.exe", b"123"))


def test_file_size_limit(monkeypatch):
    monkeypatch.setattr(document_processor, "MAX_UPLOAD_BYTES", 3)
    with pytest.raises(DocumentProcessingError, match="upload limit"):
        extract_document_text(Upload("large.txt", b"four"))


def test_text_normalization_preserves_paragraphs_and_punctuation():
    assert clean_text("Head\r\n\r\n  Line   with\x00 spaces!\x01\n\n\nEnd.") == "Head\n\nLine with spaces!\n\nEnd."
